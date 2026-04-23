#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
insert_figures.py
将专利目录下 figures/ 中的图片插入到对应的 DOCX 文档中。

策略：
  自动发现 patents/patent_*/patent_*_disclosure.docx，
  从同目录下的 figures/ 读取图片，按文件名排序依次追加到文档末尾。

用法:
  python tools/insert_figures.py              # 处理所有专利
  python tools/insert_figures.py patent_06    # 只处理 patent_06
"""

import glob
import os
import re
import sys

from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH


PATENTS_DIR = os.path.join(os.path.dirname(__file__), '..', 'patents')

# 图片宽度（英寸）
IMG_WIDTH = Inches(5.5)


def set_para_spacing(para, before_pt=0, after_pt=0, line_val=240):
    """设置段落间距。"""
    pPr = para._p.get_or_add_pPr()
    spacing = pPr.find(qn('w:spacing'))
    if spacing is None:
        spacing = OxmlElement('w:spacing')
        pPr.append(spacing)
    if before_pt:
        spacing.set(qn('w:before'), str(int(before_pt * 20)))
    if after_pt:
        spacing.set(qn('w:after'), str(int(after_pt * 20)))
    spacing.set(qn('w:lineRule'), 'auto')
    spacing.set(qn('w:line'), str(line_val))


def insert_figure_at_end(doc, img_path, caption_text):
    """在文档末尾插入图片和图注。"""
    # 图片段落
    img_para = doc.add_paragraph()
    img_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(img_para, before_pt=8, after_pt=4, line_val=240)
    run = img_para.add_run()
    run.add_picture(img_path, width=IMG_WIDTH)

    # 图注段落
    caption_para = doc.add_paragraph()
    caption_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para_spacing(caption_para, before_pt=4, after_pt=8, line_val=280)
    caption_run = caption_para.add_run(caption_text)
    caption_run.font.name = '宋体'
    caption_run.font.size = Pt(10.5)
    rPr = caption_run._r.get_or_add_rPr()
    for old in rPr.findall(qn('w:rFonts')):
        rPr.remove(old)
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), '宋体')
    rFonts.set(qn('w:hAnsi'), '宋体')
    rFonts.set(qn('w:eastAsia'), '宋体')
    rFonts.set(qn('w:cs'), '宋体')
    rPr.insert(0, rFonts)

    return True


def make_caption(fig_path, fig_num):
    """从文件名生成图注文字。如 fig_1_system_flow.png → '图1 system flow'"""
    name = os.path.splitext(os.path.basename(fig_path))[0]
    # 去掉 fig_ 前缀和尾部的 _0 等后缀
    parts = name.split('_')
    # 跳过 'fig' 和数字序号部分
    desc_parts = []
    skip = True
    for p in parts:
        if skip:
            if p == 'fig' or p.isdigit():
                continue
            skip = False
        desc_parts.append(p)
    # 去掉末尾纯数字（如生成工具追加的 _0）
    while desc_parts and desc_parts[-1].isdigit():
        desc_parts.pop()
    desc = ' '.join(desc_parts) if desc_parts else name
    return f'图{fig_num} {desc}'


def discover_patents(filter_name=None):
    """发现 patents/ 下所有专利的 docx 及其 figures/ 目录。

    返回 list of (docx_path, figures_dir)。
    """
    pattern = os.path.join(os.path.abspath(PATENTS_DIR), 'patent_*', 'patent_*_disclosure.docx')
    docx_files = sorted(glob.glob(pattern))
    results = []
    for docx_path in docx_files:
        if filter_name and filter_name not in os.path.basename(os.path.dirname(docx_path)):
            continue
        patent_dir = os.path.dirname(docx_path)
        figures_dir = os.path.join(patent_dir, 'figures')
        results.append((docx_path, figures_dir))
    return results


def get_figure_files(figures_dir):
    """获取 figures/ 目录下所有图片文件，按名称排序。"""
    if not os.path.isdir(figures_dir):
        return []
    exts = ('.png', '.jpg', '.jpeg', '.bmp', '.gif')
    files = []
    for f in sorted(os.listdir(figures_dir)):
        if os.path.splitext(f)[1].lower() in exts:
            files.append(os.path.join(figures_dir, f))
    return files


def process_patent(docx_path, figures_dir):
    """处理单个专利文档，将 figures/ 下的图片按顺序插入到文档末尾。"""
    patent_name = os.path.basename(os.path.dirname(docx_path))
    print(f"\n处理: {patent_name}")

    figures = get_figure_files(figures_dir)
    if not figures:
        print(f"  无图片（{figures_dir} 不存在或为空）")
        return 0

    doc = Document(docx_path)
    inserted = 0
    for i, fig_path in enumerate(figures, 1):
        caption = make_caption(fig_path, i)
        try:
            insert_figure_at_end(doc, fig_path, caption)
            inserted += 1
            print(f"  ✓ 插入 {caption}")
        except Exception as e:
            print(f"  ✗ 插入失败 {os.path.basename(fig_path)}: {e}")

    doc.save(docx_path)
    size_kb = os.path.getsize(docx_path) / 1024
    print(f"  保存完成: {size_kb:.1f} KB, 插入 {inserted} 张图片")
    return inserted


def main():
    import argparse
    parser = argparse.ArgumentParser(description='专利 DOCX 插图工具')
    parser.add_argument('filter', nargs='?', default=None,
                        help='可选：只处理匹配的专利，如 patent_06')
    args = parser.parse_args()

    print("=" * 60)
    print("专利 DOCX 插图插入工具")
    print("=" * 60)

    patents = discover_patents(args.filter)
    if not patents:
        print("  未发现任何 patent_*_disclosure.docx 文件。")
        print(f"  搜索路径: {os.path.abspath(PATENTS_DIR)}/patent_*/")
        sys.exit(1)

    print(f"  发现 {len(patents)} 个专利文档")

    total = 0
    for docx_path, figures_dir in patents:
        count = process_patent(docx_path, figures_dir)
        total += count

    print()
    print("=" * 60)
    print(f"完成！共插入 {total} 张插图到 {len(patents)} 个专利文档中。")
    print("=" * 60)


if __name__ == '__main__':
    main()
